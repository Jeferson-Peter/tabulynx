from __future__ import annotations

import hashlib
import uuid
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

import polars as pl
from openpyxl import load_workbook


SUPPORTED_SUFFIXES = {".csv", ".xlsx", ".parquet"}
TABULYNX_CACHE_DIR = Path.home() / ".tabulynx" / "cache"
EXCEL_CACHE_VERSION = 1
EXCEL_CHUNK_SIZE = 5_000


@dataclass(slots=True)
class DatasetPage:
    offset: int
    limit: int
    row_count: int
    columns: list[str]
    rows: list[dict[str, Any]]


@dataclass(slots=True)
class ColumnStats:
    column: str
    type_name: str
    null_count: int
    non_null_count: int
    distinct_count: int
    min_value: Any = None
    max_value: Any = None
    top_values: list[tuple[str, int]] | None = None


@dataclass(slots=True)
class DatasetProfile:
    row_count: int
    column_count: int
    type_counts: dict[str, int]
    null_heavy_columns: list[tuple[str, int]]
    numeric_columns: list[str]
    temporal_columns: list[str]
    text_columns: list[str]
    boolean_columns: list[str]


@dataclass(slots=True)
class QueryState:
    sorts: list["SortRule"] | None = None
    visible_columns: list[str] | None = None
    filters: list["ColumnFilter"] | None = None
    search_text: str | None = None

    def to_dict(self) -> dict[str, Any]:
        return {
            "sorts": [
                {"column": item.column, "descending": item.descending}
                for item in (self.sorts or [])
            ],
            "visible_columns": list(self.visible_columns) if self.visible_columns else None,
            "filters": [
                {"column": item.column, "operator": item.operator, "value": item.value}
                for item in (self.filters or [])
            ],
            "search_text": self.search_text,
        }

    @classmethod
    def from_dict(cls, data: dict[str, Any]) -> "QueryState":
        return cls(
            sorts=[
                SortRule(
                    column=item["column"],
                    descending=bool(item.get("descending", False)),
                )
                for item in data.get("sorts", [])
            ]
            or None,
            visible_columns=list(data["visible_columns"]) if data.get("visible_columns") else None,
            filters=[
                ColumnFilter(
                    column=item["column"],
                    operator=item["operator"],
                    value=item.get("value"),
                )
                for item in data.get("filters", [])
            ]
            or None,
            search_text=data.get("search_text") or None,
        )


@dataclass(slots=True)
class ColumnFilter:
    column: str
    operator: str
    value: str | None = None

    def label(self) -> str:
        if self.operator in {"is_empty", "not_empty"}:
            return f"{self.column} {self.operator}"
        return f"{self.column} {self.operator} {self.value}"


@dataclass(slots=True)
class SortRule:
    column: str
    descending: bool = False

    def label(self, position: int | None = None) -> str:
        prefix = f"{position}. " if position is not None else ""
        return f"{prefix}{self.column} {'desc' if self.descending else 'asc'}"


class DatasetSession:
    def __init__(
        self,
        path: str | Path,
        sheet: str | None = None,
        *,
        display_name: str | None = None,
        transient: bool = False,
    ) -> None:
        self.path = Path(path).expanduser().resolve()
        self.suffix = self.path.suffix.lower()
        if self.suffix not in SUPPORTED_SUFFIXES:
            raise ValueError(f"Unsupported file type: {self.suffix or 'unknown'}")
        if not self.path.exists():
            raise FileNotFoundError(self.path)

        self._data_source: Path | None = None
        self._uses_excel_cache = False
        self._display_name = display_name
        self._transient = transient
        self._sheet_names = self._discover_sheets()
        self._sheet = sheet or self._sheet_names[0]
        if self._sheet not in self._sheet_names:
            raise ValueError(f"Unknown sheet: {self._sheet}")
        self._query = QueryState()
        self._activate_source()

    @property
    def name(self) -> str:
        return self._display_name or self.path.name

    @property
    def sheet(self) -> str:
        return self._sheet

    @property
    def uses_excel_cache(self) -> bool:
        return self._uses_excel_cache

    @property
    def is_transient(self) -> bool:
        return self._transient

    def sheet_names(self) -> list[str]:
        return list(self._sheet_names)

    def set_sheet(self, sheet: str) -> None:
        if sheet not in self._sheet_names:
            raise ValueError(f"Unknown sheet: {sheet}")
        self._sheet = sheet
        self._activate_source()

    def query_state(self) -> QueryState:
        return QueryState(
            sorts=list(self._query.sorts) if self._query.sorts else None,
            visible_columns=list(self._query.visible_columns) if self._query.visible_columns else None,
            filters=list(self._query.filters) if self._query.filters else None,
            search_text=self._query.search_text,
        )

    def set_sort(self, column: str | None, descending: bool = False) -> None:
        if column is None:
            self._query.sorts = None
            return
        self.set_sorts([SortRule(column=column, descending=descending)])

    def clear_sort(self) -> None:
        self._query.sorts = None

    def set_sorts(self, sorts: list[SortRule] | None) -> None:
        if not sorts:
            self._query.sorts = None
            return
        unknown = [item.column for item in sorts if item.column not in self.columns(all_columns=True)]
        if unknown:
            raise ValueError(f"Unknown columns: {', '.join(sorted(set(unknown)))}")
        self._query.sorts = list(sorts)

    def add_sort(self, column: str, descending: bool = False) -> None:
        sorts = list(self._query.sorts or [])
        sorts = [item for item in sorts if item.column != column]
        sorts.append(SortRule(column=column, descending=descending))
        self.set_sorts(sorts)

    def remove_sort_at(self, index: int) -> None:
        sorts = list(self._query.sorts or [])
        if not (0 <= index < len(sorts)):
            return
        del sorts[index]
        self.set_sorts(sorts or None)

    def set_visible_columns(self, columns: list[str] | None) -> None:
        if columns is None:
            self._query.visible_columns = None
            return
        unknown = [column for column in columns if column not in self.columns(all_columns=True)]
        if unknown:
            raise ValueError(f"Unknown columns: {', '.join(unknown)}")
        self._query.visible_columns = list(columns)

    def set_search_text(self, value: str | None) -> None:
        normalized = (value or "").strip()
        self._query.search_text = normalized or None

    def clear_search_text(self) -> None:
        self._query.search_text = None

    def apply_query_state(self, state: QueryState) -> None:
        self.set_visible_columns(state.visible_columns)
        self.set_filters(state.filters)
        self.set_sorts(state.sorts)
        self.set_search_text(state.search_text)

    def set_filters(self, filters: list[ColumnFilter] | None) -> None:
        if not filters:
            self._query.filters = None
            return
        unknown = [item.column for item in filters if item.column not in self.columns(all_columns=True)]
        if unknown:
            raise ValueError(f"Unknown columns: {', '.join(sorted(set(unknown)))}")
        self._query.filters = list(filters)

    def clear_filters(self) -> None:
        self._query.filters = None

    def add_filter(self, column_filter: ColumnFilter) -> None:
        filters = list(self._query.filters or [])
        filters.append(column_filter)
        self.set_filters(filters)

    def remove_filter_at(self, index: int) -> None:
        filters = list(self._query.filters or [])
        if not (0 <= index < len(filters)):
            return
        del filters[index]
        self.set_filters(filters or None)

    def columns(self, all_columns: bool = False) -> list[str]:
        schema_columns = self._schema_frame().columns
        if all_columns or self._query.visible_columns is None:
            return schema_columns
        return [column for column in schema_columns if column in self._query.visible_columns]

    def dtypes(self) -> list[str]:
        frame = self._schema_frame()
        selected = self.columns()
        return [str(dtype) for column, dtype in zip(frame.columns, frame.dtypes, strict=False) if column in selected]

    def column_dtype(self, column: str) -> pl.DataType:
        frame = self._schema_frame()
        schema = dict(zip(frame.columns, frame.dtypes, strict=False))
        try:
            return schema[column]
        except KeyError as exc:
            raise ValueError(f"Unknown column: {column}") from exc

    def column_type_name(self, column: str) -> str:
        dtype = self.column_dtype(column)
        dtype_name = str(dtype)
        if dtype.is_integer():
            return "int"
        if dtype.is_float() or dtype.is_decimal():
            return "float"
        if dtype == pl.Date:
            return "date"
        if dtype_name.startswith("Datetime"):
            return "datetime"
        if dtype == pl.Boolean:
            return "bool"
        return "text"

    def row_count(self) -> int:
        lazy = self._query_frame()
        if lazy is not None:
            return int(lazy.select(pl.len()).collect().item())
        source = self._frame()
        source = self._apply_filters_frame(source)
        source = self._apply_search_frame(source)
        source = self._apply_sorts_frame(source)
        if self._query.visible_columns is not None:
            source = source.select(self.columns())
        return int(source.height)

    def preview(self, limit: int = 50) -> list[dict[str, Any]]:
        return self.page(0, limit).rows

    def current_view_frame(self) -> pl.DataFrame:
        lazy = self._query_frame()
        if lazy is not None:
            return lazy.collect()

        source = self._frame()
        source = self._apply_filters_frame(source)
        source = self._apply_search_frame(source)
        source = self._apply_sorts_frame(source)
        if self._query.visible_columns is not None:
            source = source.select(self.columns())
        return source

    def sql_query(self, query: str, *, table_name: str = "current") -> pl.DataFrame:
        frame = self.current_view_frame()
        ctx = pl.SQLContext({table_name: frame}, eager=True)
        return ctx.execute(query, eager=True)

    @classmethod
    def from_frame(
        cls,
        frame: pl.DataFrame,
        *,
        name: str = "SQL result",
    ) -> "DatasetSession":
        sql_dir = TABULYNX_CACHE_DIR / "sql"
        sql_dir.mkdir(parents=True, exist_ok=True)
        destination = sql_dir / f"{uuid.uuid4().hex}.parquet"
        frame.write_parquet(destination)
        return cls(destination, display_name=name, transient=True)

    def export_current_view(self, destination: str | Path) -> Path:
        destination_path = Path(destination).expanduser().resolve()
        suffix = destination_path.suffix.lower()
        frame = self.current_view_frame()

        if suffix == ".csv":
            frame.write_csv(destination_path)
        elif suffix == ".parquet":
            frame.write_parquet(destination_path)
        else:
            raise ValueError("Export supports only .csv and .parquet")

        return destination_path

    def column_stats(self, column: str, top_k: int = 5) -> ColumnStats:
        if column not in self.columns():
            raise ValueError(f"Unknown column: {column}")

        frame = self.current_view_frame()
        series = frame.get_column(column)
        type_name = self.column_type_name(column)
        null_count = int(series.null_count())
        non_null_count = int(len(series) - null_count)
        distinct_count = int(series.n_unique())

        non_null_series = series.drop_nulls()
        min_value = non_null_series.min() if len(non_null_series) else None
        max_value = non_null_series.max() if len(non_null_series) else None

        top_values: list[tuple[str, int]] = []
        if len(non_null_series):
            value_counts = (
                frame.select(pl.col(column))
                .drop_nulls()
                .group_by(column)
                .len()
                .sort("len", descending=True)
                .head(max(1, top_k))
                .iter_rows()
            )
            top_values = [(str(value), int(count)) for value, count in value_counts]

        return ColumnStats(
            column=column,
            type_name=type_name,
            null_count=null_count,
            non_null_count=non_null_count,
            distinct_count=distinct_count,
            min_value=min_value,
            max_value=max_value,
            top_values=top_values or None,
        )

    def dataset_profile(self, top_k: int = 5) -> DatasetProfile:
        frame = self.current_view_frame()
        columns = self.columns()
        row_count = int(frame.height)
        type_counts: dict[str, int] = {}
        null_heavy: list[tuple[str, int]] = []
        numeric_columns: list[str] = []
        temporal_columns: list[str] = []
        text_columns: list[str] = []
        boolean_columns: list[str] = []

        for column in columns:
            type_name = self.column_type_name(column)
            type_counts[type_name] = type_counts.get(type_name, 0) + 1
            null_count = int(frame.get_column(column).null_count())
            if null_count > 0:
                null_heavy.append((column, null_count))

            if type_name in {"int", "float"}:
                numeric_columns.append(column)
            elif type_name in {"date", "datetime"}:
                temporal_columns.append(column)
            elif type_name == "bool":
                boolean_columns.append(column)
            else:
                text_columns.append(column)

        null_heavy.sort(key=lambda item: item[1], reverse=True)
        return DatasetProfile(
            row_count=row_count,
            column_count=len(columns),
            type_counts=type_counts,
            null_heavy_columns=null_heavy[: max(1, top_k)],
            numeric_columns=numeric_columns,
            temporal_columns=temporal_columns,
            text_columns=text_columns,
            boolean_columns=boolean_columns,
        )

    def page(self, offset: int = 0, limit: int = 100) -> DatasetPage:
        offset = max(0, int(offset))
        limit = max(1, min(int(limit), 2_000))
        lazy = self._query_frame()
        selected_columns = self.columns()

        if lazy is not None:
            frame = lazy.slice(offset, limit).collect()
            row_count = int(lazy.select(pl.len()).collect().item())
        else:
            source = self._frame()
            source = self._apply_filters_frame(source)
            source = self._apply_search_frame(source)
            source = self._apply_sorts_frame(source)
            if self._query.visible_columns is not None:
                source = source.select(selected_columns)
            frame = source.slice(offset, limit)
            row_count = int(source.height)

        return DatasetPage(
            offset=offset,
            limit=limit,
            row_count=row_count,
            columns=selected_columns,
            rows=frame.to_dicts(),
        )

    def _activate_source(self) -> None:
        if self.suffix == ".xlsx":
            self._data_source = self._ensure_excel_cache(self._sheet)
            self._uses_excel_cache = True
            return
        self._data_source = self.path
        self._uses_excel_cache = False

    def _discover_sheets(self) -> list[str]:
        if self.suffix != ".xlsx":
            return ["Sheet1"]

        workbook = load_workbook(self.path, read_only=True)
        return list(workbook.sheetnames) or ["Sheet1"]

    def _schema_frame(self) -> pl.DataFrame:
        lazy = self._base_lazy_frame()
        if lazy is not None:
            return lazy.slice(0, 1).collect()
        return self._frame().head(1)

    def _frame(self) -> pl.DataFrame:
        if self.suffix == ".csv":
            return pl.read_csv(self.path)
        if self.suffix == ".parquet":
            return pl.read_parquet(self.path)
        if self.suffix == ".xlsx":
            if self._data_source is None:
                raise ValueError("Missing optimized source for workbook")
            return pl.read_parquet(self._data_source)
        raise ValueError(f"Unsupported file type: {self.suffix}")

    def _base_lazy_frame(self) -> pl.LazyFrame | None:
        if self.suffix == ".csv":
            return pl.scan_csv(self.path)
        if self.suffix == ".parquet":
            return pl.scan_parquet(self.path)
        if self.suffix == ".xlsx":
            if self._data_source is None:
                raise ValueError("Missing optimized source for workbook")
            return pl.scan_parquet(self._data_source)
        raise ValueError(f"Unsupported file type: {self.suffix}")

    def _query_frame(self) -> pl.LazyFrame | None:
        lazy = self._base_lazy_frame()
        if lazy is None:
            return None
        lazy = self._apply_filters_lazy(lazy)
        lazy = self._apply_search_lazy(lazy)
        lazy = self._apply_sorts_lazy(lazy)
        if self._query.visible_columns is not None:
            lazy = lazy.select(self._query.visible_columns)
        return lazy

    def _apply_filters_lazy(self, lazy: pl.LazyFrame) -> pl.LazyFrame:
        if not self._query.filters:
            return lazy
        for item in self._query.filters:
            lazy = lazy.filter(self._filter_expr(item))
        return lazy

    def _apply_filters_frame(self, frame: pl.DataFrame) -> pl.DataFrame:
        if not self._query.filters:
            return frame
        for item in self._query.filters:
            frame = frame.filter(self._filter_expr(item))
        return frame

    def _apply_search_lazy(self, lazy: pl.LazyFrame) -> pl.LazyFrame:
        if not self._query.search_text:
            return lazy
        expr = self._search_expr(self._query.search_text)
        if expr is None:
            return lazy
        return lazy.filter(expr)

    def _apply_search_frame(self, frame: pl.DataFrame) -> pl.DataFrame:
        if not self._query.search_text:
            return frame
        expr = self._search_expr(self._query.search_text)
        if expr is None:
            return frame
        return frame.filter(expr)

    def _apply_sorts_lazy(self, lazy: pl.LazyFrame) -> pl.LazyFrame:
        if not self._query.sorts:
            return lazy
        return lazy.sort(
            [item.column for item in self._query.sorts],
            descending=[item.descending for item in self._query.sorts],
        )

    def _apply_sorts_frame(self, frame: pl.DataFrame) -> pl.DataFrame:
        if not self._query.sorts:
            return frame
        return frame.sort(
            [item.column for item in self._query.sorts],
            descending=[item.descending for item in self._query.sorts],
        )

    def _filter_expr(self, item: ColumnFilter) -> pl.Expr:
        column = pl.col(item.column)
        value = item.value if item.value is not None else ""
        dtype = self.column_dtype(item.column)
        dtype_name = str(dtype)

        if item.operator == "is_empty":
            if dtype == pl.Boolean:
                return column.is_null()
            return column.is_null() | (column.cast(pl.String).fill_null("") == "")
        if item.operator == "not_empty":
            if dtype == pl.Boolean:
                return column.is_not_null()
            return column.is_not_null() & (column.cast(pl.String).fill_null("") != "")

        if dtype.is_integer():
            parsed = int(value)
            if item.operator == "equals":
                return column.cast(pl.Int64, strict=False) == parsed
            if item.operator == "gt":
                return column.cast(pl.Int64, strict=False) > parsed
            if item.operator == "gte":
                return column.cast(pl.Int64, strict=False) >= parsed
            if item.operator == "lt":
                return column.cast(pl.Int64, strict=False) < parsed
            if item.operator == "lte":
                return column.cast(pl.Int64, strict=False) <= parsed
            raise ValueError(f"Operator {item.operator} is not supported for integer columns")

        if dtype.is_float() or dtype.is_decimal():
            parsed = float(value)
            if item.operator == "equals":
                return column.cast(pl.Float64, strict=False) == parsed
            if item.operator == "gt":
                return column.cast(pl.Float64, strict=False) > parsed
            if item.operator == "gte":
                return column.cast(pl.Float64, strict=False) >= parsed
            if item.operator == "lt":
                return column.cast(pl.Float64, strict=False) < parsed
            if item.operator == "lte":
                return column.cast(pl.Float64, strict=False) <= parsed
            raise ValueError(f"Operator {item.operator} is not supported for float columns")

        if dtype == pl.Date:
            parsed = pl.lit(value).str.strptime(pl.Date, strict=True)
            if item.operator == "equals":
                return column.cast(pl.Date, strict=False) == parsed
            if item.operator == "gt":
                return column.cast(pl.Date, strict=False) > parsed
            if item.operator == "gte":
                return column.cast(pl.Date, strict=False) >= parsed
            if item.operator == "lt":
                return column.cast(pl.Date, strict=False) < parsed
            if item.operator == "lte":
                return column.cast(pl.Date, strict=False) <= parsed
            raise ValueError(f"Operator {item.operator} is not supported for date columns")

        if dtype_name.startswith("Datetime"):
            parsed = pl.lit(value).str.strptime(pl.Datetime, strict=True)
            casted = column.cast(pl.Datetime, strict=False)
            if item.operator == "equals":
                return casted == parsed
            if item.operator == "gt":
                return casted > parsed
            if item.operator == "gte":
                return casted >= parsed
            if item.operator == "lt":
                return casted < parsed
            if item.operator == "lte":
                return casted <= parsed
            raise ValueError(f"Operator {item.operator} is not supported for datetime columns")

        if dtype == pl.Boolean:
            normalized = value.lower()
            if normalized not in {"true", "false", "1", "0", "yes", "no"}:
                raise ValueError("Boolean filters accept: true/false, 1/0, yes/no")
            parsed = normalized in {"true", "1", "yes"}
            if item.operator == "equals":
                return column.cast(pl.Boolean, strict=False) == parsed
            raise ValueError(f"Operator {item.operator} is not supported for boolean columns")

        if item.operator == "contains":
            return column.cast(pl.String).str.contains(value, literal=True, strict=False)
        if item.operator == "equals":
            return column.cast(pl.String) == value
        if item.operator == "starts_with":
            return column.cast(pl.String).str.starts_with(value)
        if item.operator == "ends_with":
            return column.cast(pl.String).str.ends_with(value)
        raise ValueError(f"Unsupported filter operator: {item.operator}")

    def _search_expr(self, value: str) -> pl.Expr | None:
        columns = self.columns(all_columns=True)
        if not columns:
            return None

        expressions = [
            pl.col(column).cast(pl.String, strict=False).fill_null("").str.contains(value, literal=True, strict=False)
            for column in columns
        ]
        return pl.any_horizontal(expressions)

    def _ensure_excel_cache(self, sheet: str) -> Path:
        cache_dir = self._excel_cache_dir(sheet)
        cache_file = cache_dir / "data.parquet"
        if cache_file.exists():
            return cache_file

        cache_dir.mkdir(parents=True, exist_ok=True)
        headers, schema_map = self._infer_excel_headers_and_schema(sheet)
        self._write_excel_sheet_to_parquet(sheet=sheet, headers=headers, schema_map=schema_map, destination=cache_file)
        return cache_file

    def _excel_cache_dir(self, sheet: str) -> Path:
        stat = self.path.stat()
        key = hashlib.sha256(
            f"{self.path}|{stat.st_size}|{stat.st_mtime_ns}|{sheet}|{EXCEL_CACHE_VERSION}".encode("utf-8")
        ).hexdigest()[:24]
        return TABULYNX_CACHE_DIR / key

    def _infer_excel_headers_and_schema(self, sheet: str) -> tuple[list[str], dict[str, pl.DataType]]:
        workbook = load_workbook(self.path, read_only=True, data_only=True)
        try:
            worksheet = workbook[sheet]
            rows = worksheet.iter_rows(values_only=True)
            first_row = next(rows, None)
            headers = self._normalize_excel_headers(first_row)
            kinds = ["empty"] * len(headers)
            for row in rows:
                normalized = self._normalize_excel_row(row, len(headers))
                for index, value in enumerate(normalized):
                    kinds[index] = self._merge_excel_kind(kinds[index], self._excel_value_kind(value))
        finally:
            workbook.close()

        schema_map = {
            header: self._polars_dtype_for_kind(kind)
            for header, kind in zip(headers, kinds, strict=False)
        }
        return headers, schema_map

    def _write_excel_sheet_to_parquet(
        self,
        sheet: str,
        headers: list[str],
        schema_map: dict[str, pl.DataType],
        destination: Path,
    ) -> None:
        workbook = load_workbook(self.path, read_only=True, data_only=True)
        chunks: list[pl.DataFrame] = []
        try:
            worksheet = workbook[sheet]
            rows = worksheet.iter_rows(values_only=True)
            next(rows, None)
            buffer: list[list[Any]] = []
            for row in rows:
                buffer.append(self._normalize_excel_row(row, len(headers)))
                if len(buffer) >= EXCEL_CHUNK_SIZE:
                    chunks.append(self._excel_chunk_frame(buffer, headers, schema_map))
                    buffer = []

            if buffer:
                chunks.append(self._excel_chunk_frame(buffer, headers, schema_map))

            if chunks:
                frame = pl.concat(chunks, how="vertical_relaxed")
            else:
                frame = pl.DataFrame({header: [] for header in headers}).cast(schema_map)
            frame.write_parquet(destination)
        finally:
            workbook.close()

    def _excel_chunk_frame(
        self,
        rows: list[list[Any]],
        headers: list[str],
        schema_map: dict[str, pl.DataType],
    ) -> pl.DataFrame:
        columns: dict[str, list[Any]] = {header: [] for header in headers}
        for row in rows:
            for index, header in enumerate(headers):
                dtype = schema_map[header]
                columns[header].append(self._coerce_excel_value(row[index], dtype))
        return pl.DataFrame(columns, schema=schema_map)

    def _normalize_excel_headers(self, first_row: tuple[Any, ...] | None) -> list[str]:
        raw_headers = list(first_row or [])
        if not raw_headers:
            return ["column_1"]

        headers: list[str] = []
        seen: dict[str, int] = {}
        for index, value in enumerate(raw_headers, start=1):
            base = str(value).strip() if value not in (None, "") else f"column_{index}"
            count = seen.get(base, 0)
            seen[base] = count + 1
            headers.append(base if count == 0 else f"{base}_{count + 1}")
        return headers

    def _normalize_excel_row(self, row: tuple[Any, ...] | None, width: int) -> list[Any]:
        values = list(row or [])
        if len(values) < width:
            values.extend([None] * (width - len(values)))
        return values[:width]

    def _excel_value_kind(self, value: Any) -> str:
        if value is None:
            return "empty"
        if isinstance(value, bool):
            return "bool"
        if isinstance(value, datetime):
            return "datetime"
        if isinstance(value, date):
            return "date"
        if isinstance(value, time):
            return "time"
        if isinstance(value, int):
            return "int"
        if isinstance(value, float):
            return "float"
        return "text"

    def _merge_excel_kind(self, current: str, new: str) -> str:
        if new == "empty":
            return current
        if current == "empty":
            return new
        if current == new:
            return current
        if {current, new} <= {"int", "float"}:
            return "float"
        if {current, new} <= {"date", "datetime"}:
            return "datetime"
        return "text"

    def _polars_dtype_for_kind(self, kind: str) -> pl.DataType:
        if kind == "bool":
            return pl.Boolean
        if kind == "int":
            return pl.Int64
        if kind == "float":
            return pl.Float64
        if kind == "date":
            return pl.Date
        if kind == "datetime":
            return pl.Datetime
        return pl.String

    def _coerce_excel_value(self, value: Any, dtype: pl.DataType) -> Any:
        if value is None:
            return None
        if dtype == pl.String:
            return str(value)
        if dtype == pl.Boolean:
            return bool(value) if not isinstance(value, str) else value.lower() in {"true", "1", "yes"}
        if dtype == pl.Int64:
            return int(value) if value != "" else None
        if dtype == pl.Float64:
            return float(value) if value != "" else None
        if dtype == pl.Date:
            if isinstance(value, datetime):
                return value.date()
            if isinstance(value, date):
                return value
            return None
        if dtype == pl.Datetime:
            if isinstance(value, datetime):
                return value
            if isinstance(value, date):
                return datetime.combine(value, time.min)
            return None
        return value


def open_dataset(path: str | Path, sheet: str | None = None) -> DatasetSession:
    return DatasetSession(path=path, sheet=sheet)
